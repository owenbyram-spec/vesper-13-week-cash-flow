"""Reopen exported artifacts and reconcile them to immutable exact forecasts."""

from copy import deepcopy
import csv
from decimal import Decimal
from io import BytesIO, StringIO
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from vesper_cashflow.exports import csv_export, excel_export, printable_export
from vesper_cashflow.imports import synthetic_demo
from vesper_cashflow.proposals import delay_receipts
from vesper_cashflow.store import Store


def exported_version(tmp_path, inputs=None):
    inputs = inputs or synthetic_demo()
    store = Store(tmp_path)
    workspace = store.create_workspace("Synthetic export review", currency=inputs["currency"], synthetic=True)
    version = store.save_version(inputs, "Base", scope=workspace)
    return store, store.active_workspace(), version


def assert_numeric_money(cell, exact, places=2):
    assert cell.data_type == "n"
    assert Decimal(str(cell.value)) == Decimal(exact)
    assert Decimal.from_float(float(cell.value)).quantize(Decimal(1).scaleb(-places)) == Decimal(exact)
    assert cell.number_format.startswith("#,##0" + ("." + "0" * places if places else ";"))


def test_excel_numeric_forecast_and_inputs_reconcile_with_selected_versions(tmp_path):
    store, workspace, base = exported_version(tmp_path)
    proposal = store.draft_proposal(
        delay_receipts(base["inputs"], 14), "Delayed collections", base_hash=base["input_hash"]
    )
    approved = store.resolve_proposal(proposal["id"], proposal["proposal_hash"], True)["version"]
    for version in (base, approved):
        workbook = load_workbook(BytesIO(excel_export(workspace, version)))
        result = version["result"]
        forecast = workbook["Forecast"]
        assert forecast.max_row == 14
        for row, week in enumerate(result["weeks"], start=2):
            assert forecast.cell(row, 1).value == week["number"]
            assert forecast.cell(row, 1).data_type == "n"
            assert forecast.cell(row, 2).value == week["start"]
            assert forecast.cell(row, 3).value == week["end"]
            for column, field in enumerate(("beginning_cash", "cash_in", "cash_out", "ending_cash"), start=4):
                assert_numeric_money(forecast.cell(row, column), week[field])
            amounts = [Decimal(str(forecast.cell(row, column).value)) for column in range(4, 8)]
            assert amounts[0] + amounts[1] - amounts[2] == amounts[3]
        assert sum(Decimal(str(forecast.cell(row, 5).value)) for row in range(2, 15)) == Decimal(result["total_in"])
        assert sum(Decimal(str(forecast.cell(row, 6).value)) for row in range(2, 15)) == Decimal(result["total_out"])
        for sheet_name, records, column in (
            ("Cash items", version["inputs"]["items"], 5),
            ("Recurring per occurrence", version["inputs"]["recurrences"], 4),
            ("Dated occurrences", result["occurrences"], 5),
        ):
            sheet = workbook[sheet_name]
            assert sheet.max_row == len(records) + 1
            for row, record in enumerate(records, start=2):
                assert_numeric_money(sheet.cell(row, column), record["amount"])
        metadata = {row[0].value: row[1] for row in workbook["Version"].iter_rows(min_row=2)}
        assert metadata["Version"].value == version["id"]
        assert metadata["Input SHA-256"].value == version["input_hash"]
        assert_numeric_money(metadata["Opening cash"], version["inputs"]["opening_cash"])
        assert_numeric_money(metadata["Minimum cash"], result["minimum_cash"]["amount"])
        assert metadata["Sequence"].value == version["sequence"]
        assert metadata["Coverage complete"].value is result["coverage_complete"]
        assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    assert store.get_version(base["id"])["result"] == base["result"]
    assert approved["result"]["minimum_cash"]["amount"] == "15500.00"


@pytest.mark.parametrize("currency,value,places", [("USD", "-25.55", 2), ("KWD", "-25.123", 3), ("JPY", "-25", 0)])
def test_excel_money_scale_negative_and_zero_are_numeric(tmp_path, currency, value, places):
    inputs = synthetic_demo()
    inputs["currency"] = currency
    inputs["opening_cash"] = value
    for record in inputs["items"] + inputs["recurrences"]:
        record["currency"] = currency
    _, workspace, version = exported_version(tmp_path, inputs)
    workbook = load_workbook(BytesIO(excel_export(workspace, version)))
    assert_numeric_money(workbook["Forecast"]["D2"], value, places)
    assert_numeric_money(workbook["Forecast"]["E3"], "0", places)


def test_excel_precision_limit_keeps_exact_text_and_explains_it(tmp_path):
    inputs = synthetic_demo()
    inputs["opening_cash"] = "999999999999998.99"
    inputs["items"][0]["amount"] = "10000000000000.01"
    _, workspace, version = exported_version(tmp_path, inputs)
    workbook = load_workbook(BytesIO(excel_export(workspace, version)))
    opening = workbook["Forecast"]["D2"]
    assert opening.data_type == "s"
    assert opening.value == inputs["opening_cash"]
    assert "15 significant digits" in opening.comment.text
    assert workbook["Cash items"]["E2"].value == inputs["items"][0]["amount"]
    assert workbook["Cash items"]["E2"].data_type == "s"
    assert_numeric_money(workbook["Cash items"]["E3"], inputs["items"][1]["amount"])
    assert inputs["opening_cash"] in csv_export(workspace, version).decode("utf-8-sig")


@pytest.mark.parametrize("value", ["99999999999999.90", "99999999999999.80", "9999999999999.99"])
def test_excel_serialized_xml_preserves_high_cent_amounts(tmp_path, value):
    inputs = synthetic_demo()
    inputs["opening_cash"] = value
    _, workspace, version = exported_version(tmp_path, inputs)
    content = excel_export(workspace, version)
    workbook = load_workbook(BytesIO(content))
    cell = workbook["Forecast"]["D2"]
    with ZipFile(BytesIO(content)) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet2.xml"))
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    xml_cell = root.find(".//s:c[@r='D2']", ns)
    if value == "99999999999999.90":
        # Although only 15 significant decimal digits remain after normalization,
        # this value becomes .91 in numeric XLSX XML and must stay exact text.
        assert cell.data_type == "s"
        assert cell.value == value
        assert cell.comment is not None
        assert xml_cell.find("s:is/s:t", ns).text == value
    else:
        assert_numeric_money(cell, value)
        assert Decimal(xml_cell.find("s:v", ns).text) == Decimal(value)


def test_export_text_safety_ids_and_exact_csv_print_rows(tmp_path):
    inputs = deepcopy(synthetic_demo())
    inputs["items"][0]["id"] = "000012345678901234"
    inputs["items"][0]["name"] = '=HYPERLINK("unsafe")'
    inputs["items"][0]["source"] = "+CMD()"
    _, workspace, version = exported_version(tmp_path, inputs)
    workbook = load_workbook(BytesIO(excel_export(workspace, version)))
    assert workbook["Cash items"]["A2"].value == "000012345678901234"
    assert workbook["Cash items"]["A2"].data_type == "s"
    assert workbook["Cash items"]["B2"].value.startswith("'=")
    assert workbook["Cash items"]["J2"].value == "'+CMD()"
    rows = list(csv.reader(StringIO(csv_export(workspace, version).decode("utf-8-sig"))))
    forecast_index = rows.index(["Forecast"])
    headers = rows[forecast_index + 1]
    for exported, week in zip(rows[forecast_index + 2:forecast_index + 15], version["result"]["weeks"]):
        assert exported == [str(week[key]) for key in headers]
    html = printable_export(workspace, version)
    assert version["id"] in html
    for week in version["result"]["weeks"]:
        assert "<tr>" + "".join(f"<td>{week[key]}</td>" for key in headers) + "</tr>" in html
