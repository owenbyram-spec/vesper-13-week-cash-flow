"""Exact selected-version exports with provenance and spreadsheet text safety."""

from __future__ import annotations

import csv
import io
import re
from decimal import Decimal
from html import escape

from .domain.money import Money


def safe_cell(value: object) -> str:
    value = "" if value is None else str(value)
    if re.fullmatch(r"-?\d+(?:\.\d+)?", value):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def tables(workspace: dict, version: dict) -> list[tuple[str, list[str], list[list]]]:
    inputs, result = version["inputs"], version["result"]
    metadata = [
        ["Entity", workspace["name"]],
        ["Workspace", workspace["id"]],
        ["Synthetic", "YES — fictional teaching example" if workspace["synthetic"] else "No"],
        ["Version", version["id"]],
        ["Sequence", version["sequence"]],
        ["Label", version["label"]],
        ["Saved", version["created_at"]],
        ["As of (closing cash)", inputs["as_of"]],
        ["Currency", inputs["currency"]],
        ["Opening cash", inputs["opening_cash"]],
        ["Opening source", inputs["opening_source"]],
        ["Input SHA-256", version["input_hash"]],
        ["Result SHA-256", version["result_hash"]],
        ["Algorithm", result["algorithm_version"]],
        ["Coverage complete", result["coverage_complete"]],
        ["Minimum cash", result["minimum_cash"]["amount"]],
        ["Minimum observed on", result["minimum_cash"]["date"]],
        ["Minimum observation basis", result["minimum_cash"]["basis"]],
        ["Notes", inputs["notes"]],
    ]

    def rows(entries, keys):
        return [[entry.get(key, "") for key in keys] for entry in entries]

    weekly_keys = ["number", "start", "end", "beginning_cash", "cash_in", "cash_out", "ending_cash"]
    item_keys = [
        "id",
        "name",
        "category",
        "direction",
        "amount",
        "currency",
        "due_date",
        "expected_date",
        "timing_basis",
        "source",
        "source_id",
        "ordinal",
    ]
    recurrence_keys = [
        "id",
        "name",
        "direction",
        "amount",
        "currency",
        "cadence",
        "anchor_date",
        "start_date",
        "end_date",
        "timing_basis",
        "source",
    ]
    occurrence_keys = [
        "id",
        "recurrence_id",
        "name",
        "direction",
        "amount",
        "currency",
        "expected_date",
        "timing_basis",
        "source",
    ]
    source_keys = ["name", "kind", "sha256", "observed_at", "details"]
    return [
        ("Version", ["Field", "Value"], metadata),
        ("Forecast", weekly_keys, rows(result["weeks"], weekly_keys)),
        ("Cash items", item_keys, rows(inputs["items"], item_keys)),
        ("Recurring amounts per occurrence", recurrence_keys, rows(inputs["recurrences"], recurrence_keys)),
        ("Dated occurrences", occurrence_keys, rows(result["occurrences"], occurrence_keys)),
        ("Coverage", ["Section", "Status"], [[key, value] for key, value in inputs["coverage"].items()]),
        (
            "Findings",
            ["Code", "Severity", "Items", "Message"],
            [
                [entry["code"], entry["severity"], ", ".join(entry["ids"]), entry["message"]]
                for entry in result["findings"]
            ],
        ),
        ("Sources", source_keys, rows(inputs["sources"], source_keys)),
    ]


def csv_export(workspace: dict, version: dict) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    for title, headers, rows in tables(workspace, version):
        writer.writerow([safe_cell(title)])
        writer.writerow(headers)
        writer.writerows([[safe_cell(value) for value in row] for row in rows])
        writer.writerow([])
    return output.getvalue().encode("utf-8-sig")


def excel_export(workspace: dict, version: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    places = Money.MINOR_UNITS.get(version["inputs"]["currency"], 2)
    quantum = Decimal(1).scaleb(-places)
    money_format = "#,##0" + ("." + "0" * places if places else "")
    money_format += ";[Red](" + money_format + ")"
    money_headers = {"amount", "beginning_cash", "cash_in", "cash_out", "ending_cash"}
    names = [
        "Version",
        "Forecast",
        "Cash items",
        "Recurring per occurrence",
        "Dated occurrences",
        "Coverage",
        "Findings",
        "Sources",
    ]
    for sheet_name, (_, headers, rows) in zip(names, tables(workspace, version)):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        for row_number, row in enumerate(rows, start=2):
            for column, value in enumerate(row, start=1):
                cell = sheet.cell(row_number, column)
                is_money = headers[column - 1] in money_headers or (
                    sheet_name == "Version" and column == 2 and row[0] in {"Opening cash", "Minimum cash"}
                )
                if is_money:
                    exact = Decimal(value)
                    # Both Excel's 15-digit precision and the XLSX writer's
                    # binary conversion must preserve the amount at currency scale.
                    # str(float) alone hides cent changes near the precision limit.
                    numeric_safe = (
                        len(exact.normalize().as_tuple().digits) <= 15
                        and Decimal.from_float(float(exact)).quantize(quantum) == exact
                        and Decimal("%.16g" % exact) == exact
                    )
                    if numeric_safe:
                        cell.value = exact
                        cell.number_format = money_format
                    else:
                        cell.value = safe_cell(value)
                        cell.comment = Comment(
                            "Exact amount stored as text to preserve currency precision beyond Excel's "
                            "15 significant digits or safe binary numeric conversion. "
                            "Use the raw CSV or JSON backup for exact decimal processing.",
                            "Vesper & Co",
                        )
                elif headers[column - 1] in {"number", "ordinal"} or (
                    sheet_name == "Version" and column == 2 and row[0] in {"Sequence", "Coverage complete"}
                ):
                    cell.value = value
                else:
                    # IDs, dates and user text stay text, including formula-like text.
                    cell.value = safe_cell(value)
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="142D3D")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(65, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def printable_export(workspace: dict, version: dict) -> str:
    sections = []
    for title, headers, rows in tables(workspace, version):
        header = "".join(f"<th>{escape(value)}</th>" for value in headers)
        body = "".join(
            "<tr>"
            + "".join(f"<td>{escape(str(value if value is not None else ''))}</td>" for value in row)
            + "</tr>"
            for row in rows
        )
        sections.append(
            f"<section><h2>{escape(title)}</h2><table><thead><tr>{header}</tr></thead><tbody>{body}</tbody></table></section>"
        )
    synthetic = "<p class='notice'>SYNTHETIC — fictional teaching data</p>" if workspace["synthetic"] else ""
    return (
        "<!doctype html><html><head><meta charset='utf-8'><title>Vesper &amp; Co · 13-week cash flow</title><style>body{font:12px system-ui;color:#142d3d;margin:32px}h1,h2{font-family:Georgia,serif}h1{font-size:30px}h2{margin-top:32px}table{border-collapse:collapse;width:100%;table-layout:auto}td,th{border:1px solid #d8dad4;padding:6px;text-align:left;overflow-wrap:anywhere}th{background:#f2eee6}tr{break-inside:avoid}section{break-inside:auto}.notice{background:#f2e3d3;padding:12px}@media print{body{margin:0;font-size:9px}h2{break-after:avoid}@page{size:landscape;margin:12mm}}</style></head><body><h1>Vesper &amp; Co</h1><p>13-week cash flow · Use your browser’s Print → Save as PDF.</p>"
        + synthetic
        + "".join(sections)
        + "</body></html>"
    )
