"""Local boundaries around the unchanged native engine, with fictional inputs.

Each case owns a fresh process because native settings and SQLAlchemy are global.
No fixture uses a provider key, OS secret store, live accounting source or legacy DB.
"""

from pathlib import Path
import subprocess
import sys

import pytest


@pytest.mark.parametrize(
    "scenario",
    [
        "forecast_export",
        "boundaries",
        "money",
        "actuals",
        "partial_actuals",
        "workspace_keys",
        "qbo_binding",
        "arap",
        "refresh_replay",
        "finite_spread",
        "storage_warnings",
        "scenarios_variance",
        "partial_refresh",
    ],
)
def test_local_native_backend(scenario):
    result = subprocess.run(
        [sys.executable, str(Path(__file__).resolve()), scenario],
        cwd=Path(__file__).resolve().parents[2],
        capture_output=True,
        text=True,
        timeout=50,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def _exercise(scenario):
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    import io
    import json
    import os
    import tempfile
    import time
    from datetime import date

    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from pe_platform.backend import bootstrap, cashflow

    # Ambient credentials must never become the local workspace connection.
    os.environ["OPENAI_API_KEY"] = "fixture-never-a-provider-key"
    directory = Path(tempfile.mkdtemp(prefix="vesper-native-backend-test-"))
    runtime = bootstrap.configure(directory, port=8893, testing=True)
    from app.config import settings
    from app.tenant_rls import tenant_scope

    assert not settings.openai_api_key and not settings.anthropic_api_key
    app = FastAPI()
    bootstrap.register(app)
    if scenario == "finite_spread":
        from datetime import timedelta
        from app.routers import cashflow_forecast as native

        # Include a Sunday monthly anchor and a year boundary deterministically.
        native._week_fridays = lambda n=13, lookback=0: [
            date(2026, 10, 30) + timedelta(weeks=index) for index in range(-lookback, n)
        ]
    with TestClient(app) as client:
        initial = client.get("/local/bootstrap")
        assert initial.status_code == 200, initial.text
        headers = {"x-csrf-token": initial.json()["csrf_token"]}
        response = client.post("/local/demo", headers=headers)
        assert response.status_code == 200, response.text
        scope = response.json()["workspace"]
        headers.update({"x-vesper-workspace": scope["id"], "x-vesper-revision": str(scope["revision"])})
        params = {"company_id": scope["realm_id"]}
        response = client.get("/cashflow-forecast/weekly", params=params, headers=headers)
        assert response.status_code == 200, response.text
        grid = response.json()

        if scenario == "forecast_export":
            from openpyxl import load_workbook
            from sqlalchemy.exc import IntegrityError

            assert len(grid["week_headers"]) == 14 and grid["lookback"] == 1
            assert all(date.fromisoformat(day).weekday() == 4 for day in grid["week_headers"])
            assert not grid["data_quality"]["is_partial"]
            assert grid["bank_balance"] == 50000
            for index, ending in enumerate(grid["ending_balances"]):
                assert (
                    round(
                        grid["beginning_balances"][index]
                        + grid["total_cash_in"][index]
                        - grid["total_cash_out"][index],
                        2,
                    )
                    == ending
                )
            fingerprint = cashflow.financial_fingerprint(runtime, scope)
            first = cashflow.capture_export_snapshot(runtime, scope)
            second = cashflow.capture_export_snapshot(runtime, scope)
            assert first["export_id"] == second["export_id"]
            assert cashflow.financial_fingerprint(runtime, scope) == fingerprint
            export_id = first["export_id"]
            for operation in (
                "UPDATE cash_forecast_export_snapshots SET content_hash='changed'",
                "DELETE FROM cash_forecast_export_snapshots",
            ):
                try:
                    with runtime.engine.begin() as db:
                        db.exec_driver_sql(operation)
                except IntegrityError as exc:
                    assert "immutable" in str(exc)
                else:
                    raise AssertionError("Native immutable export content was writable")
            path = f"/reporting/cashflow-forecast/exports/{export_id}/download"
            download = client.get(
                path, params={"expected_company_id": scope["realm_id"], "format": "excel"}, headers=headers
            )
            assert download.status_code == 200, download.text
            workbook = load_workbook(io.BytesIO(download.content), data_only=True)
            rows = [row for sheet in workbook for row in sheet.iter_rows(values_only=True)]
            ending_row = next(row for row in rows if "Ending Bank Balance" in row)
            label_index = ending_row.index("Ending Bank Balance")
            numeric = [value for value in ending_row[label_index + 1 :] if isinstance(value, (int, float))]
            assert numeric[-14:] == grid["ending_balances"], ending_row
            # Later financial edits cannot change an exact-ID download.
            before = json.dumps(first, sort_keys=True)
            result = client.post(
                "/cashflow-forecast/override",
                json={
                    "company_id": scope["realm_id"],
                    "entity_name": "Other Cash In",
                    "entity_type": "customer",
                    "week_ending": grid["week_headers"][1],
                    "override_amount": 123.45,
                },
                headers=headers,
            )
            assert result.status_code == 200, result.text
            assert cashflow.financial_fingerprint(runtime, scope) != fingerprint
            from app.cashflow.export_snapshots import load_forecast_export_from_authorized_row
            from app.models import CashForecastExportSnapshot
            from uuid import UUID

            with (
                tenant_scope(org_id=scope["organization_id"], realm_id=scope["realm_id"]),
                runtime.session_factory() as db,
            ):
                row = (
                    db.query(CashForecastExportSnapshot)
                    .filter(CashForecastExportSnapshot.id == UUID(export_id))
                    .one()
                )
                preserved = load_forecast_export_from_authorized_row(
                    row,
                    expected_company_id=scope["realm_id"],
                    expected_organization_id=scope["organization_id"],
                ).model_dump(mode="json")
            assert json.dumps(preserved, sort_keys=True) == before

        elif scenario == "boundaries":
            assert client.get("/cashflow-forecast/weekly", params=params).status_code == 409
            assert (
                client.get(
                    "/cashflow-forecast/weekly",
                    params=params,
                    headers={**headers, "host": "attacker.example"},
                ).status_code
                == 403
            )
            assert (
                client.get(
                    "/cashflow-forecast/weekly",
                    params=params,
                    headers={**headers, "origin": "https://attacker.example"},
                ).status_code
                == 403
            )
            assert (
                client.get(
                    "/cashflow-forecast/weekly", params={"company_id": "another-company"}, headers=headers
                ).status_code
                == 403
            )
            assert (
                client.get(
                    "/cashflow-forecast/actuals/pull",
                    params=params,
                    headers={key: value for key, value in headers.items() if key != "x-csrf-token"},
                ).status_code
                == 403
            )
            assert (
                client.post(
                    "/local/demo",
                    headers={key: value for key, value in headers.items() if key != "x-csrf-token"},
                ).status_code
                == 403
            )
            result = client.post("/local/workspaces", headers=headers, json={"name": "Second local entity"})
            assert result.status_code == 200, result.text
            second = result.json()["workspace"]
            assert second["organization_id"] != scope["organization_id"]
            assert second["user_id"] != scope["user_id"]
            assert client.get("/cashflow-forecast/weekly", params=params, headers=headers).status_code == 409
            second_headers = {
                **headers,
                "x-vesper-workspace": second["id"],
                "x-vesper-revision": str(second["revision"]),
            }
            old_export = f"/reporting/cashflow-forecast/exports/{grid['export_id']}/download"
            denied = client.get(
                old_export,
                params={"expected_company_id": second["realm_id"], "format": "excel"},
                headers=second_headers,
            )
            assert denied.status_code == 404, denied.text
            client.cookies.clear()
            assert client.get("/cashflow-forecast/weekly", params=params, headers=headers).status_code == 401

        elif scenario == "money":
            from pe_platform.backend.security import validate_money_tree

            validate_money_tree(
                {"weekly_amounts": [{"week_ending": grid["week_headers"][1], "amount": 75.25}]}
            )
            for invalid in (0.001, "NaN", "Infinity", "1000000000000.00", True):
                result = client.post(
                    "/cashflow-forecast/override",
                    json={
                        "company_id": scope["realm_id"],
                        "entity_name": "Other Cash In",
                        "entity_type": "customer",
                        "week_ending": grid["week_headers"][1],
                        "override_amount": invalid,
                    },
                    headers=headers,
                )
                assert result.status_code == 400, result.text
            result = client.post(
                "/local/workspaces", json={"name": "Non-USD", "currency": "GBP"}, headers=headers
            )
            assert result.status_code == 400 and "USD only" in result.text

        elif scenario in {"actuals", "partial_actuals"}:
            from pe_platform.backend.fixtures import SyntheticQBOClient
            from app.models import CashForecastWeekStatus

            if scenario == "partial_actuals":
                original = SyntheticQBOClient.query

                def fail_one_source(self, query):
                    if "FROM Purchase " in query:
                        raise RuntimeError("Fictional unavailable source")
                    return original(self, query)

                SyntheticQBOClient.query = fail_one_source
            target = grid["week_headers"][0]
            result = client.get(
                "/cashflow-forecast/actuals/pull", params={**params, "week_ending": target}, headers=headers
            )
            assert result.status_code == 200, result.text
            with (
                tenant_scope(org_id=scope["organization_id"], realm_id=scope["realm_id"]),
                runtime.session_factory() as db,
            ):
                status = (
                    db.query(CashForecastWeekStatus)
                    .filter(
                        CashForecastWeekStatus.company_id == scope["realm_id"],
                        CashForecastWeekStatus.week_ending == date.fromisoformat(target),
                    )
                    .one()
                )
                assert bool(status.pull_partial) == (scenario == "partial_actuals")
            result = client.post(
                "/cashflow-forecast/actuals/close",
                params={**params, "week_ending": target},
                json={},
                headers=headers,
            )
            assert result.status_code == (409 if scenario == "partial_actuals" else 200), result.text
            if scenario == "actuals":
                assert result.json()["status"] == "closed"

        elif scenario == "workspace_keys":
            status = client.get("/local/connections", headers=headers).json()
            assert status["model"]["configured"] is False
            saved = client.post(
                "/local/connections/model",
                json={
                    "provider": "openai",
                    "api_key": "fictional-session-value",
                    "model": "fixture-model",
                    "persist": False,
                    "consent": True,
                },
                headers=headers,
            )
            assert saved.status_code == 200, saved.text
            assert "fictional-session-value" not in saved.text
            assert client.get("/local/connections", headers=headers).json()["model"]["configured"] is True
            second = client.post("/local/workspaces", headers=headers, json={"name": "No model here"}).json()[
                "workspace"
            ]
            second_headers = {
                **headers,
                "x-vesper-workspace": second["id"],
                "x-vesper-revision": str(second["revision"]),
            }
            assert (
                client.get("/local/connections", headers=second_headers).json()["model"]["configured"]
                is False
            )
            assert runtime.settings.get(scope["id"], "model")["api_key"] == "fictional-session-value"
            assert not settings.openai_api_key

        elif scenario == "storage_warnings":
            from vesper_cashflow.settings import SettingsError

            class FakePersistentStore:
                persistent = True
                locked = False
                values = {}

                def get(self, key):
                    if self.locked:
                        raise SettingsError(
                            "The OS credential store could not be read; unlock it or reconnect for this session."
                        )
                    return self.values.get(key)

                def set(self, key, value):
                    if self.locked:
                        raise SettingsError(
                            "The OS credential store could not save the credentials; choose session-only storage."
                        )
                    self.values[key] = value

                def delete(self, key):
                    if self.locked:
                        raise SettingsError(
                            "The stored credentials could not be removed; unlock the OS credential store and retry."
                        )
                    self.values.pop(key, None)

            store = FakePersistentStore()
            runtime.settings._store = store
            body = {
                "provider": "openai",
                "api_key": "fictional-persistent-value",
                "model": "fixture-model",
                "persist": True,
                "consent": True,
            }
            result = client.post("/local/connections/model", json=body, headers=headers)
            assert result.status_code == 200 and result.json()["storage_mode"] == "os-credential-store"
            key = runtime.settings._key(scope["id"], "model")
            original_copy = store.values[key]
            runtime.settings.session.delete(key)
            runtime.settings.storage_modes.pop(key)
            store.locked = True
            status = client.get("/local/connections", headers=headers).json()
            assert not status["model"]["configured"] and "could not be read" in status["storage_warning"]
            result = client.post("/local/connections/model", json=body, headers=headers)
            assert result.status_code == 400 and store.values[key] == original_copy
            result = client.post(
                "/local/connections/model",
                json={**body, "api_key": "fictional-session-value", "persist": False},
                headers=headers,
            )
            assert result.status_code == 200
            status = client.get("/local/connections", headers=headers).json()
            assert status["model"]["configured"] and status["model"]["storage_mode"] == "session-only"
            assert "were not checked or removed" in status["storage_warning"]
            assert "fictional-" not in json.dumps(status)
            assert store.values[key] == original_copy
            assert client.delete("/local/connections/model", headers=headers).status_code == 400
            assert runtime.settings.get(scope["id"], "model") is not None
            store.locked = False
            assert client.delete("/local/connections/model", headers=headers).status_code == 200
            assert key not in store.values and runtime.settings.get(scope["id"], "model") is None

        elif scenario == "qbo_binding":
            second = client.post(
                "/local/workspaces", headers=headers, json={"name": "Bound local company"}
            ).json()["workspace"]
            with tenant_scope(org_id=second["organization_id"], realm_id=second["realm_id"]):
                bound = runtime.workspaces.bind_realm(second, "1234567890")
            assert bound["realm_id"] == "1234567890" and bound["revision"] > second["revision"]
            from pe_platform.backend.connections import NativeReadOnlyClient
            from pe_platform.backend.fixtures import SyntheticQBOClient, dataset
            import httpx

            read_client = NativeReadOnlyClient(runtime, bound)
            for report in ("../BalanceSheet", "GeneralLedger", "BalanceSheet?x=y"):
                try:
                    read_client.get_report(report)
                except Exception as exc:
                    assert "fixed cash-flow" in str(exc)
                else:
                    raise AssertionError("Arbitrary report paths allowed")
            try:
                read_client.query("DELETE FROM Invoice")
            except Exception as exc:
                assert "SELECT" in str(exc)
            else:
                raise AssertionError("Accounting write query allowed")
            # Exercise production GET handling with injected HTTP bytes, not a
            # real QBO authorization or acquisition claim.
            calls = []
            fixture = SyntheticQBOClient(bound, dataset())
            bundle = {
                "realm_id": bound["realm_id"],
                "environment": "sandbox",
                "access_token": "fictional-transport-value",
            }
            read_client.legacy._require_bound = lambda: None
            read_client.legacy._refresh = lambda **_kwargs: bundle

            def fixture_http(method, url, **kwargs):
                calls.append((method, url, kwargs["params"]))
                assert method == "GET" and url.startswith(
                    "https://sandbox-quickbooks.api.intuit.com/v3/company/1234567890/"
                )
                report = url.rsplit("/", 1)[-1]
                payload = (
                    fixture.get_report(report, kwargs["params"])
                    if "/reports/" in url
                    else fixture.query(kwargs["params"]["query"])
                )
                return httpx.Response(200, json=payload)

            read_client.legacy._http = fixture_http
            for report in (
                "BalanceSheet",
                "AgedReceivables",
                "AgedReceivableDetail",
                "AgedPayables",
                "AgedPayableDetail",
                "ProfitAndLoss",
            ):
                assert read_client.get_report(report)["Header"]["Currency"] == "USD"
            assert read_client.query("SELECT * FROM Invoice MAXRESULTS 1000")["QueryResponse"]["Invoice"]
            assert len(calls) == 7 and all(call[2]["minorversion"] == "73" for call in calls)
            foreign_currency = fixture.get_report("BalanceSheet")
            foreign_currency["Header"]["Currency"] = "GBP"
            read_client.legacy._http = lambda *_args, **_kwargs: httpx.Response(200, json=foreign_currency)
            try:
                read_client.get_report("BalanceSheet")
            except Exception as exc:
                assert "USD currency" in str(exc)
            else:
                raise AssertionError("Foreign accounting currency silently aggregated")

        elif scenario == "arap":
            from app.services import ar_ap_summary

            @app.get("/test-arap")
            def get_summary():
                with runtime.session_factory() as db:
                    return ar_ap_summary.serve_summary(db, scope["realm_id"])

            result = client.get("/test-arap", headers=headers)
            assert result.status_code == 200, result.text
            assert result.json()["data_quality"] == "computing"
            for _ in range(100):
                result = client.get("/test-arap", headers=headers)
                if result.json()["data_quality"] != "computing":
                    break
                time.sleep(0.01)
            payload = result.json()
            assert payload["receivables"]["total"] == 12000, payload
            assert payload["payables"]["total"] == 6000
            assert payload["annual_revenue"] == 48000 and payload["annual_cogs"] == 18000
            assert payload["data_quality"]["is_partial"] is False

        elif scenario == "refresh_replay":
            from pe_platform.backend.connections import replay_source
            from pe_platform.backend.fixtures import SyntheticQBOClient
            from app.models import CashForecastOverride

            request = {
                "company_id": scope["realm_id"],
                "entity_name": "Other Cash In",
                "entity_type": "customer",
                "week_ending": grid["week_headers"][1],
                "override_amount": 123.45,
            }
            result = client.post("/cashflow-forecast/override", json=request, headers=headers)
            assert result.status_code == 200, result.text
            result = client.post("/cashflow-forecast/weekly/refresh", params=params, json={}, headers=headers)
            assert result.status_code == 200, result.text
            with (
                tenant_scope(org_id=scope["organization_id"], realm_id=scope["realm_id"]),
                runtime.session_factory() as db,
            ):
                saved = (
                    db.query(CashForecastOverride)
                    .filter(CashForecastOverride.company_id == scope["realm_id"])
                    .all()
                )
                assert any(float(row.override_amount) == 123.45 for row in saved)
            source_before = runtime.workspaces.source_path(scope).read_bytes()
            fingerprint = cashflow.financial_fingerprint(runtime, scope)
            original = SyntheticQBOClient.query

            def no_acquisition(*_args, **_kwargs):
                raise AssertionError("Replay attempted source acquisition")

            SyntheticQBOClient.query = no_acquisition
            with runtime.session_factory() as db, replay_source(scope):
                replayed = cashflow.cached_forecast(runtime, scope, db)
                assert replayed["week_headers"] == grid["week_headers"]
            assert runtime.workspaces.source_path(scope).read_bytes() == source_before
            assert cashflow.financial_fingerprint(runtime, scope) == fingerprint
            SyntheticQBOClient.query = original

        elif scenario == "finite_spread":
            base_request = {
                "company_id": scope["realm_id"],
                "entity_name": "Synthetic finite weekly",
                "entity_type": "vendor",
                "total_amount": 1200,
                "periods": 3,
                "cadence": "weekly",
                "start_week_ending": "2026-10-30",
                "day_of_month": 15,
            }
            response = client.post("/cashflow-forecast/spread-balance", json=base_request, headers=headers)
            assert response.status_code == 200, response.text
            assert response.json()["avg_amount"] == 0
            monthly_request = {
                **base_request,
                "entity_name": "Synthetic finite monthly",
                "total_amount": 1000,
                "cadence": "monthly",
            }
            response = client.post("/cashflow-forecast/spread-balance", json=monthly_request, headers=headers)
            assert response.status_code == 200, response.text
            assert response.json()["weekly_amounts"] == [
                {"week_ending": "2026-10-30", "amount": 333.33},
                {"week_ending": "2026-11-20", "amount": 333.33},
                {"week_ending": "2026-12-18", "amount": 333.34},
            ]

            def assert_finite_grid():
                current = client.get("/cashflow-forecast/weekly", params=params, headers=headers).json()
                weekly = next(row for row in current["vendors"] if row["name"] == base_request["entity_name"])
                monthly = next(
                    row for row in current["vendors"] if row["name"] == monthly_request["entity_name"]
                )
                actual_weekly = {
                    week: amount for week, amount in zip(current["week_headers"], weekly["amounts"]) if amount
                }
                actual_monthly = {
                    week: amount
                    for week, amount in zip(current["week_headers"], monthly["amounts"])
                    if amount
                }
                assert actual_weekly == {"2026-10-30": 400, "2026-11-06": 400, "2026-11-13": 400}, (
                    actual_weekly
                )
                assert actual_monthly == {
                    item["week_ending"]: item["amount"] for item in response.json()["weekly_amounts"]
                }, actual_monthly
                assert sum(weekly["amounts"]) == 1200 and round(sum(monthly["amounts"]), 2) == 1000

            assert_finite_grid()
            result = client.post("/cashflow-forecast/weekly/refresh", params=params, json={}, headers=headers)
            assert result.status_code == 200, result.text
            assert_finite_grid()
            fingerprint = cashflow.financial_fingerprint(runtime, scope)
            collision = {**base_request, "entity_name": grid["vendors"][0]["name"]}
            assert (
                client.post("/cashflow-forecast/spread-balance", json=collision, headers=headers).status_code
                == 409
            )
            assert (
                client.post(
                    "/cashflow-forecast/spread-balance", json=base_request, headers=headers
                ).status_code
                == 409
            )
            for invalid in (
                {"periods": 50},
                {"periods": True},
                {"day_of_month": 0},
                {"day_of_month": 29},
                {"cadence": "monthly", "periods": 4, "day_of_month": 28},
            ):
                result = client.post(
                    "/cashflow-forecast/spread-balance",
                    json={**base_request, "entity_name": "Never created", **invalid},
                    headers=headers,
                )
                assert result.status_code == 400, result.text
            assert cashflow.financial_fingerprint(runtime, scope) == fingerprint

        elif scenario == "scenarios_variance":
            from app.cashflow.scenarios import apply_scenario_transform
            from app.cashflow.variance import build_week_variance_report
            from app.cashflow.versioning import actuals_by_entity_for_weeks, compact_grid_from_payload

            base_json = json.dumps(grid, sort_keys=True)
            fingerprint = cashflow.financial_fingerprint(runtime, scope)
            first = client.get("/cashflow-forecast/weekly/scenarios", params=params, headers=headers)
            second = client.get("/cashflow-forecast/weekly/scenarios", params=params, headers=headers)
            assert first.status_code == second.status_code == 200
            scenarios = first.json()
            assert scenarios["writeback"] == "read_only"
            assert scenarios["cases"] == second.json()["cases"]
            for key in (
                "week_headers",
                "total_cash_in",
                "total_cash_out",
                "beginning_balances",
                "ending_balances",
            ):
                assert scenarios["base"][key] == grid[key]
            assert {case["case"] for case in scenarios["cases"]} == {
                "downside_collections",
                "revenue_shock",
                "cost_inflation",
            }
            delay = apply_scenario_transform(
                grid, name="synthetic_delay", label="Synthetic 14-day delay", ar_slip_weeks=2
            )
            upside = apply_scenario_transform(
                grid, name="synthetic_upside", label="Synthetic receipts +10%", collections_haircut_pct=-0.10
            )
            downside = next(case for case in scenarios["cases"] if case["case"] == "cost_inflation")
            lookback = grid["lookback"]
            assert delay["total_cash_in"][:lookback] == grid["total_cash_in"][:lookback]
            assert delay["total_cash_in"][lookback : lookback + 2] == [0, 0]
            assert delay["total_cash_in"][lookback + 2 :] == grid["total_cash_in"][lookback:-2]
            assert round(
                sum(delay["total_cash_in"][lookback:]) + delay["slipped_beyond_horizon"], 2
            ) == round(sum(grid["total_cash_in"][lookback:]), 2)
            for index in range(lookback, len(grid["week_headers"])):
                assert upside["total_cash_in"][index] == round(grid["total_cash_in"][index] * 1.1, 2)
                assert downside["total_cash_out"][index] == round(grid["total_cash_out"][index] * 1.1, 2)
            for case in (delay, upside, downside):
                assert case["beginning_balances"][lookback] == grid["beginning_balances"][lookback]
                assert case["total_cash_out"][:lookback] == grid["total_cash_out"][:lookback]
            assert json.dumps(grid, sort_keys=True) == base_json
            assert cashflow.financial_fingerprint(runtime, scope) == fingerprint
            target = grid["week_headers"][0]
            result = client.get(
                "/cashflow-forecast/weekly/variance",
                params={**params, "week_ending": target},
                headers=headers,
            )
            assert result.status_code == 200 and result.json()["status"] == "no_frozen_version"
            result = client.get(
                "/cashflow-forecast/actuals/pull", params={**params, "week_ending": target}, headers=headers
            )
            assert result.status_code == 200 and result.json()["count"] == 3, result.text
            with (
                tenant_scope(org_id=scope["organization_id"], realm_id=scope["realm_id"]),
                runtime.session_factory() as db,
            ):
                captured_actuals = actuals_by_entity_for_weeks(
                    db, scope["realm_id"], [date.fromisoformat(target)]
                )[date.fromisoformat(target)]
            report = build_week_variance_report(compact_grid_from_payload(grid), target, captured_actuals)
            assert report["status"] == "scored"
            assert report["totals"]["actual_in"] == 2000 and report["totals"]["actual_out"] == 1500
            assert report["totals"]["actual_in"] - report["totals"]["actual_out"] == 500
            for row in report["rows"]:
                assert row["variance"] == round(row["actual"] - row["forecast"], 2)
            # Helper math over a captured synthetic grid is not backfilled history.
            result = client.get(
                "/cashflow-forecast/weekly/variance",
                params={**params, "week_ending": target},
                headers=headers,
            )
            assert result.json()["status"] == "no_frozen_version"

        elif scenario == "partial_refresh":
            from pe_platform.backend.fixtures import SyntheticQBOClient
            from fastapi import HTTPException

            request = {
                "company_id": scope["realm_id"],
                "entity_name": "Other Cash In",
                "entity_type": "customer",
                "week_ending": grid["week_headers"][1],
                "override_amount": 123.45,
            }
            assert (
                client.post("/cashflow-forecast/override", json=request, headers=headers).status_code == 200
            )
            complete_fingerprint = cashflow.financial_fingerprint(runtime, scope)
            original = SyntheticQBOClient.query

            def unavailable_history(self, query):
                if "FROM Purchase " in query:
                    raise RuntimeError("Fictional history unavailable")
                return original(self, query)

            SyntheticQBOClient.query = unavailable_history
            result = client.post("/cashflow-forecast/weekly/refresh", params=params, json={}, headers=headers)
            assert result.status_code == 200, result.text
            assert result.json()["data_quality"]["is_partial"] is True
            assert result.json()["data_quality"]["unavailable_sections"]
            assert result.json()["other_cash_in"][1] == 123.45
            assert cashflow.financial_fingerprint(runtime, scope) != complete_fingerprint
            with runtime.session_factory() as db:
                try:
                    cashflow.cached_forecast(runtime, scope, db)
                except HTTPException as exc:
                    assert exc.status_code == 409 and "incomplete" in exc.detail
                else:
                    raise AssertionError("Missing source coverage admitted an AI proposal preview")
            SyntheticQBOClient.query = original
            result = client.post("/cashflow-forecast/weekly/refresh", params=params, json={}, headers=headers)
            assert result.status_code == 200 and not result.json()["data_quality"]["is_partial"]
            assert result.json()["other_cash_in"][1] == 123.45

    assert not (directory / "forecast.sqlite3").exists()
    try:
        bootstrap.configure(directory / "different", port=8893, testing=True)
    except RuntimeError:
        pass
    else:
        raise AssertionError("Native global DB silently repointed")


if __name__ == "__main__":
    _exercise(sys.argv[1])
