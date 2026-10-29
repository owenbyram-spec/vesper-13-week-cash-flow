"""
app/fpa/uaf/excel_ingest.py
───────────────────────────
Flexible tabular ingestor for Trial Balances from Excel (.xlsx) and CSV files.
Supports arbitrary column naming conventions, signed balances, debit/credit pairs,
and auto-detects header offsets.

Amount cells are read STRICTLY (:func:`parse_amount`): only the declared
conventions below are accepted, and every other cell is reported back as a
cell-level error — an unreadable amount never becomes 0.0 or a different number.
The trial balance period is always stated by the caller (``period_end``, plus
``period_start`` whenever the file carries income-statement balances); it is
never defaulted to the upload date.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple, Union

try:
    import openpyxl
except ImportError:
    openpyxl = None  # Handled gracefully if xlsx parsing invoked without openpyxl

from .schema import (
    NormalBalance,
    UniversalAccount,
    UniversalClassification,
    UniversalTrialBalance,
)


def _clean_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


class TrialBalanceParseError(ValueError):
    """One or more amount cells fall outside the declared conventions.

    ``cell_errors`` lists every offending cell as {row, column, value, reason}
    (``row`` is the 1-based line in the source file) so the user can fix the
    file; nothing is ingested while any cell is unreadable.
    """

    def __init__(self, cell_errors: List[Dict[str, Any]]):
        self.cell_errors = cell_errors
        first = "; ".join(
            f"row {e['row']} {e['column']}: {e['reason']}" for e in cell_errors[:5]
        )
        more = f" (+{len(cell_errors) - 5} more)" if len(cell_errors) > 5 else ""
        super().__init__(f"{len(cell_errors)} trial-balance amount cell(s) could not be read: {first}{more}")


class TrialBalancePeriodError(ValueError):
    """The stated trial-balance period is missing or inconsistent."""

    def __init__(self, field_name: str, message: str):
        self.field_name = field_name
        super().__init__(message)


_CENT = Decimal("0.01")
# The schema stores floats: below 10 trillion a cent amount (15 significant
# digits) survives the float round-trip exactly; above it a cell is refused.
_MAX_AMOUNT = Decimal("10000000000000")
# Number bodies: plain digits or comma-grouped thousands, at most two decimals.
_AMOUNT_BODY = re.compile(r"(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d{1,2})?|\.\d{1,2}")
_DR_CR_SUFFIX = re.compile(r"^(.*?)\s*(dr|cr)\.?$", re.IGNORECASE)


def parse_amount(val: Any) -> Tuple[Optional[Decimal], Optional[str]]:
    """Strictly parse ONE trial-balance amount cell.

    Returns ``(amount, side)``: ``amount`` is a cent-exact Decimal (``None`` for a
    blank cell) and ``side`` is ``"dr"``/``"cr"`` when the cell carried a Dr/Cr
    suffix, else ``None``. Raises ``ValueError(reason)`` for anything else.

    Declared conventions (text cells): an optional ``$``; a negative written as a
    leading ``-`` (``-1,000.00``, ``-$1,000.00``, ``$-1,000.00``), parentheses
    (``(1,000.00)``, ``($1,000.00)``, ``$(1,000.00)``) or a trailing ``-``
    (``1,000.00-``); comma thousands separators in groups of three; at most two
    decimal places; an optional ``Dr``/``Cr`` suffix (never combined with a sign);
    and a lone ``-`` as the accounting-format zero. European decimals
    (``1.234,56``), scientific notation, extra decimals and any other text are
    refused. Numeric cells (Excel) must be finite and cent-exact. Any amount of
    10 trillion or more is refused (``_MAX_AMOUNT``).
    """
    if val is None:
        return None, None
    if isinstance(val, bool):
        raise ValueError("not an amount")
    if isinstance(val, (int, float, Decimal)):
        try:
            d = Decimal(str(val))
        except InvalidOperation as exc:
            raise ValueError("not a finite number") from exc
        if not d.is_finite():
            raise ValueError("not a finite number")
        if abs(d) >= _MAX_AMOUNT:
            raise ValueError("amount too large")
        q = d.quantize(_CENT)
        if abs(d - q) > Decimal("0.000001"):
            raise ValueError("more than two decimal places")
        return q, None
    if not isinstance(val, str):
        raise ValueError("unrecognized amount format")

    s = val.strip()
    if not s:
        return None, None
    side = None
    m = _DR_CR_SUFFIX.match(s)
    if m:
        s, side = m.group(1).strip(), m.group(2).lower()
    negative = False
    if s.startswith("$"):
        s = s[1:].strip()
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1].strip()
    elif s.startswith("-"):
        negative, s = True, s[1:].strip()
    elif s.endswith("-") and len(s) > 1:
        negative, s = True, s[:-1].strip()
    if s.startswith("$"):
        s = s[1:].strip()
    if s.startswith("(") or s.startswith("-") or s.endswith("-"):
        raise ValueError("ambiguous sign")
    if s == "" and negative and side is None and val.strip().lstrip("$").strip() == "-":
        return Decimal("0.00"), None  # the accounting-format zero
    if not _AMOUNT_BODY.fullmatch(s):
        raise ValueError("unrecognized amount format")
    if negative and side is not None:
        raise ValueError("ambiguous sign: a Dr/Cr suffix together with a minus sign or parentheses")
    amount = Decimal(s.replace(",", ""))
    if amount >= _MAX_AMOUNT:
        raise ValueError("amount too large")
    amount = amount.quantize(_CENT)
    return (-amount if negative else amount), side


def _infer_classification(
    name: str,
    account_num: Optional[str] = None,
    type_hint: Optional[str] = None,
) -> Tuple[UniversalClassification, NormalBalance]:
    """Infer financial statement classification and normal balance side."""
    hint = (type_hint or "").lower().strip()

    if hint in ("revenue", "income", "sales"):
        return UniversalClassification.REVENUE, NormalBalance.CREDIT
    if hint in ("cogs", "cost of goods sold", "cost of sales"):
        return UniversalClassification.COGS, NormalBalance.DEBIT
    if hint in ("expense", "operating expense", "opex", "sga"):
        return UniversalClassification.OPEX, NormalBalance.DEBIT
    if hint in ("current asset", "bank", "cash", "ar", "receivable"):
        return UniversalClassification.CURRENT_ASSET, NormalBalance.DEBIT
    if hint in ("fixed asset", "non-current asset", "ppe", "intangible"):
        return UniversalClassification.NON_CURRENT_ASSET, NormalBalance.DEBIT
    if hint in ("current liability", "ap", "payable", "credit card"):
        return UniversalClassification.CURRENT_LIABILITY, NormalBalance.CREDIT
    if hint in ("long term liability", "debt", "notes payable"):
        return UniversalClassification.NON_CURRENT_LIABILITY, NormalBalance.CREDIT
    if hint in ("equity", "retained earnings", "capital"):
        return UniversalClassification.EQUITY, NormalBalance.CREDIT

    # Infer by account number prefix if available
    num_str = (account_num or "").strip()
    if num_str and num_str[0].isdigit():
        prefix = num_str[0]
        if prefix == "1":
            return UniversalClassification.CURRENT_ASSET, NormalBalance.DEBIT
        elif prefix == "2":
            return UniversalClassification.CURRENT_LIABILITY, NormalBalance.CREDIT
        elif prefix == "3":
            return UniversalClassification.EQUITY, NormalBalance.CREDIT
        elif prefix == "4":
            return UniversalClassification.REVENUE, NormalBalance.CREDIT
        elif prefix == "5":
            return UniversalClassification.COGS, NormalBalance.DEBIT
        elif prefix in ("6", "7", "8"):
            return UniversalClassification.OPEX, NormalBalance.DEBIT
        elif prefix == "9":
            return UniversalClassification.OTHER_EXPENSE, NormalBalance.DEBIT

    # Infer by name keywords
    n = name.lower()
    if any(k in n for k in ["revenue", "sales", "fee income", "billings", "subscription"]):
        return UniversalClassification.REVENUE, NormalBalance.CREDIT
    if any(k in n for k in ["cost of goods", "cost of sales", "cogs", "direct labor", "materials"]):
        return UniversalClassification.COGS, NormalBalance.DEBIT
    if any(k in n for k in ["payroll", "salary", "rent", "insurance", "software", "utilities", "marketing", "travel", "legal", "accounting", "office"]):
        return UniversalClassification.OPEX, NormalBalance.DEBIT
    if any(k in n for k in ["cash", "checking", "savings", "accounts receivable", "prepaid", "inventory"]):
        return UniversalClassification.CURRENT_ASSET, NormalBalance.DEBIT
    if any(k in n for k in ["equipment", "machinery", "leasehold", "goodwill", "patents", "depreciation"]):
        return UniversalClassification.NON_CURRENT_ASSET, NormalBalance.DEBIT
    if any(k in n for k in ["accounts payable", "accrued", "credit card", "payroll tax payable", "current portion"]):
        return UniversalClassification.CURRENT_LIABILITY, NormalBalance.CREDIT
    if any(k in n for k in ["term loan", "subordinated note", "line of credit", "mortgage", "notes payable"]):
        return UniversalClassification.NON_CURRENT_LIABILITY, NormalBalance.CREDIT
    if any(k in n for k in ["common stock", "preferred stock", "retained earnings", "member equity", "paid-in capital", "distributions"]):
        return UniversalClassification.EQUITY, NormalBalance.CREDIT

    return UniversalClassification.UNKNOWN, NormalBalance.DEBIT


def _match_column(
    headers: List[str],
    candidates: List[str],
    exclude: Optional[set] = None,
) -> Optional[int]:
    """Find index of header matching any candidate name.

    Matches in tiers (exact header, then whole token, then bare substring) and
    skips indices already claimed by another role, so a decorated header like
    "YTD Debit" or "Ending Debit Balance" cannot also win the balance slot.
    """
    skip = exclude or set()
    norm_headers = [re.sub(r"[^\w]", "", h.lower()) for h in headers]
    header_tokens = [set(re.findall(r"[a-z0-9]+", h.lower())) for h in headers]
    norm_candidates = [
        (re.sub(r"[^\w]", "", c.lower()), c.lower().split()) for c in candidates
    ]

    for tier in ("exact", "token", "substring"):
        for norm_c, words in norm_candidates:
            if not norm_c:
                continue
            for idx, h in enumerate(norm_headers):
                if idx in skip:
                    continue
                if tier == "exact" and norm_c == h:
                    return idx
                if tier == "token" and words and all(w in header_tokens[idx] for w in words):
                    return idx
                if tier == "substring" and norm_c in h:
                    return idx
    return None


# Headers that name a side/kind rather than carrying an amount.
_NON_AMOUNT_HEADERS = {
    "normalbalance", "normalside", "balancetype", "balanceside", "balancesheet",
    "accounttype", "drcr", "debitcredit", "sign", "side",
}

_SUMMARY_LABELS = {
    "total", "totals", "subtotal", "sub total", "grand total", "sum", "summary",
    "net income", "net loss", "net income (loss)", "net change", "net",
    "discrepancy", "difference", "out of balance",
}

_SUMMARY_SECTIONS = {
    "assets", "current assets", "non current assets", "noncurrent assets",
    "fixed assets", "other assets", "liabilities", "current liabilities",
    "long term liabilities", "non current liabilities", "liabilities and equity",
    "equity", "revenue", "revenues", "income", "expense", "expenses",
    "operating expenses", "cogs", "cost of goods sold", "debits", "credits",
}


def _is_summary_label(label: str) -> bool:
    """True only for genuine total/summary rows, never for a real account that
    merely contains "sum"/"total" (e.g. "Consumer Sales", "Total Rewards Payable")."""
    norm = re.sub(r"\s+", " ", label.lower().strip().strip(":-")).strip()
    if not norm:
        return False
    if norm in _SUMMARY_LABELS:
        return True
    for lead in ("total", "totals", "subtotal", "grand total", "sum of"):
        if norm.startswith(lead + " ") and norm[len(lead) + 1:] in _SUMMARY_SECTIONS:
            return True
    return False


def _normalize_net_signs(
    entries: List[Tuple[UniversalAccount, Decimal, Decimal, Optional[Decimal]]],
    debit_positive: bool = False,
) -> List[Tuple[UniversalAccount, Decimal, Decimal, Optional[Decimal]]]:
    """Normalize signed net columns to normal-side-positive.

    Exports commonly use the debit-positive/credit-negative convention; downstream
    totals assume every net is positive on its account's normal side. A balance
    column written with Dr/Cr suffixes STATES that convention (``debit_positive``),
    so it converts without the majority inference.
    """
    credit_nets = [
        net for acct, _d, _c, net in entries
        if net and acct.normal_balance == NormalBalance.CREDIT
    ]
    if not credit_nets:
        return entries
    negatives = sum(1 for n in credit_nets if n < 0)
    if not debit_positive and negatives * 2 <= len(credit_nets):
        return entries
    return [
        (
            acct,
            debit,
            credit,
            -net if net is not None and acct.normal_balance == NormalBalance.CREDIT else net,
        )
        for acct, debit, credit, net in entries
    ]


def parse_tabular_rows(
    rows: List[List[Any]],
    entity_id: str,
    *,
    period_end: date,
    period_start: Optional[date] = None,
    currency: str = "USD",
) -> UniversalTrialBalance:
    """
    Parse generic row grid (from CSV or Excel) into UniversalTrialBalance.
    Automatically detects header row, column roles, and data rows.

    ``period_end`` is the trial balance date; ``period_start`` is required when any
    account is (or may be) an income-statement account, because those balances
    cover a period rather than a point in time. Raises
    :class:`TrialBalanceParseError` listing every unreadable amount cell and
    :class:`TrialBalancePeriodError` for a missing or inverted period.
    """
    if not rows:
        raise ValueError("Cannot ingest empty tabular dataset")
    if not isinstance(period_end, date):
        raise TrialBalancePeriodError("period_end", "State the trial balance period end date.")
    if period_start is not None and period_start > period_end:
        raise TrialBalancePeriodError("period_start", "The period start must not be after the period end.")

    as_of = period_end

    # Step 1: Detect header row (first row with at least 2 non-empty string cells matching common TB terms)
    header_idx = -1
    col_mapping: Dict[str, int] = {}

    for idx, row in enumerate(rows[:15]):  # Look in first 15 rows
        row_str = [_clean_str(c) for c in row]
        if not any(row_str):
            continue

        acct_idx = _match_column(row_str, ["account", "acct", "code", "account number", "account code"])
        name_idx = _match_column(row_str, ["name", "account name", "description", "title", "account description"])
        debit_idx = _match_column(row_str, ["debit", "dr", "debits"])
        credit_idx = _match_column(row_str, ["credit", "cr", "credits"])
        type_idx = _match_column(row_str, ["type", "classification", "category", "account type"])
        # A metadata column ("Normal Balance", "Balance Type") holds a side name, not an
        # amount; letting it win the balance slot parses every net as 0.0.
        non_amount = {
            i for i, h in enumerate(row_str)
            if re.sub(r"[^\w]", "", h.lower()) in _NON_AMOUNT_HEADERS
        }
        balance_idx = _match_column(
            row_str,
            ["balance", "net balance", "amount", "ending balance", "ytd"],
            exclude=non_amount | {
                i for i in (debit_idx, credit_idx, type_idx) if i is not None
            },
        )

        # We need either (acct or name) AND (debit+credit OR balance)
        has_id = (acct_idx is not None) or (name_idx is not None)
        has_amount = (debit_idx is not None and credit_idx is not None) or (balance_idx is not None)

        if has_id and has_amount:
            header_idx = idx
            if acct_idx is not None:
                col_mapping["account_num"] = acct_idx
            if name_idx is not None:
                col_mapping["name"] = name_idx
            if debit_idx is not None:
                col_mapping["debit"] = debit_idx
            if credit_idx is not None:
                col_mapping["credit"] = credit_idx
            if balance_idx is not None:
                col_mapping["balance"] = balance_idx
            if type_idx is not None:
                col_mapping["type"] = type_idx
            break

    if header_idx == -1:
        raise ValueError("Could not detect valid Trial Balance header columns (requires Account/Name and Debit/Credit or Balance)")

    # Fallbacks if name or account_num is missing
    if "name" not in col_mapping and "account_num" in col_mapping:
        col_mapping["name"] = col_mapping["account_num"]
    elif "account_num" not in col_mapping and "name" in col_mapping:
        col_mapping["account_num"] = col_mapping["name"]

    tb = UniversalTrialBalance(
        entity_id=entity_id,
        as_of_date=as_of,
        currency=currency,
        source_metadata={"header_row": header_idx, "column_mapping": col_mapping},
    )

    header_row = rows[header_idx]

    def _header(col: int) -> str:
        return _clean_str(header_row[col]) if col < len(header_row) else f"column {col + 1}"

    # Step 2: Iterate data rows
    pending: List[Tuple[UniversalAccount, Decimal, Decimal, Optional[Decimal]]] = []
    skipped: List[str] = []
    cell_errors: List[Dict[str, Any]] = []
    balance_sides: List[Tuple[int, Optional[str], Optional[Decimal]]] = []
    for r_idx, row in enumerate(rows[header_idx + 1:], start=header_idx + 1):
        if not row or not any(_clean_str(c) for c in row):
            continue

        raw_name = _clean_str(row[col_mapping["name"]]) if col_mapping.get("name") is not None and col_mapping["name"] < len(row) else ""
        raw_num = _clean_str(row[col_mapping["account_num"]]) if col_mapping.get("account_num") is not None and col_mapping["account_num"] < len(row) else ""

        # Skip total / summary rows, recording each so the loss is never silent
        label = raw_name or raw_num
        if _is_summary_label(label):
            skipped.append(label)
            continue

        if not raw_name and not raw_num:
            continue

        account_id = raw_num or raw_name or f"acct_{r_idx}"
        display_name = raw_name or raw_num

        def _error(col: int, value: Any, reason: str) -> None:
            cell_errors.append({
                "row": r_idx + 1, "column": _header(col),
                "value": _clean_str(value)[:40], "reason": reason,
            })

        def _cell(role: str) -> Tuple[Optional[Decimal], Optional[str], bool]:
            col = col_mapping.get(role)
            if col is None or col >= len(row):
                return None, None, True
            try:
                amount, side = parse_amount(row[col])
            except ValueError as exc:
                _error(col, row[col], str(exc))
                return None, None, False
            return amount, side, True

        debit, debit_side, ok_d = _cell("debit")
        credit, credit_side, ok_c = _cell("credit")
        net, net_side, ok_n = _cell("balance")
        if debit_side == "cr":
            _error(col_mapping["debit"], row[col_mapping["debit"]], "Dr/Cr suffix contradicts the column")
            ok_d = False
        if credit_side == "dr":
            _error(col_mapping["credit"], row[col_mapping["credit"]], "Dr/Cr suffix contradicts the column")
            ok_c = False
        if not (ok_d and ok_c and ok_n):
            continue
        if net is not None and net_side == "cr":
            net = -net
        if "balance" in col_mapping:
            balance_sides.append((r_idx, net_side, net))

        if debit is None and credit is None and net is None:
            if re.sub(r"\s+", " ", label.lower().strip().strip(":-")).strip() in _SUMMARY_SECTIONS:
                skipped.append(label)  # a bare section heading carries no amount
                continue
            col = next(col_mapping[r] for r in ("balance", "debit", "credit") if r in col_mapping)
            _error(col, "", "blank amount (enter 0 for a zero balance)")
            continue

        type_hint = ""
        if "type" in col_mapping and col_mapping["type"] < len(row):
            type_hint = _clean_str(row[col_mapping["type"]])

        classification, normal_side = _infer_classification(
            name=display_name,
            account_num=raw_num,
            type_hint=type_hint,
        )

        account = UniversalAccount(
            account_id=account_id,
            name=display_name,
            account_number=raw_num if raw_num else None,
            classification=classification,
            sub_type=type_hint if type_hint else None,
            normal_balance=normal_side,
            source_system="tabular_import",
        )

        pending.append((account, debit or Decimal("0.00"), credit or Decimal("0.00"), net))

    # A balance column states Dr/Cr on every nonzero amount or on none: a mixed
    # column leaves the side of each unsuffixed amount unknown.
    suffixed = any(side is not None for _r, side, _n in balance_sides)
    if suffixed:
        col = col_mapping["balance"]
        for r_idx, side, net in balance_sides:
            if side is None and net:
                cell_errors.append({
                    "row": r_idx + 1, "column": _header(col),
                    "value": _clean_str(rows[r_idx][col])[:40],
                    "reason": "the balance column mixes Dr/Cr-suffixed and unsuffixed amounts",
                })

    if cell_errors:
        raise TrialBalanceParseError(cell_errors)
    if not pending:
        raise ValueError("No data rows could be extracted from Trial Balance")

    if period_start is None and any(not acct.is_balance_sheet() for acct, _d, _c, _n in pending):
        raise TrialBalancePeriodError(
            "period_start",
            "The trial balance carries income-statement (or unclassified) balances, which cover a "
            "period: state the period start date as well as the period end.",
        )

    for account, debit, credit, net in _normalize_net_signs(pending, debit_positive=suffixed):
        # The schema stores floats; cent-exact Decimals convert without drift.
        tb.add_entry(
            account=account,
            debit=float(debit),
            credit=float(credit),
            net=float(net) if net is not None else None,
            period_start=None if account.is_balance_sheet() else period_start,
            period_end=period_end,
        )

    tb.source_metadata["skipped_summary_rows"] = skipped
    tb.source_metadata["skipped_row_count"] = len(skipped)
    tb.source_metadata["period_start"] = period_start.isoformat() if period_start else None
    tb.source_metadata["period_end"] = period_end.isoformat()

    return tb


def ingest_csv_trial_balance(
    csv_content: Union[str, bytes],
    entity_id: str,
    *,
    period_end: date,
    period_start: Optional[date] = None,
    currency: str = "USD",
) -> UniversalTrialBalance:
    """Ingest a CSV trial balance string or byte content for a stated period."""
    if isinstance(csv_content, bytes):
        try:
            text = csv_content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("The trial balance CSV must be UTF-8 text") from exc
    else:
        text = csv_content

    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    return parse_tabular_rows(
        rows, entity_id, period_end=period_end, period_start=period_start, currency=currency
    )


def ingest_excel_trial_balance(
    excel_bytes: bytes,
    entity_id: str,
    sheet_name: Optional[str] = None,
    *,
    period_end: date,
    period_start: Optional[date] = None,
    currency: str = "USD",
) -> UniversalTrialBalance:
    """Ingest an Excel (.xlsx) trial balance from bytes for a stated period (the
    same strict amount parser as the CSV path)."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to parse Excel trial balances")

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))

    return parse_tabular_rows(
        rows, entity_id, period_end=period_end, period_start=period_start, currency=currency
    )
