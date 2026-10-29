# PLAN-EXEMPT: AC-1 engine decomposition under ratified plan 20260629-custom-reporting; the 500-line file cap forces the registry into context+sections helper modules.
"""
app/reporting/excel_context.py
───────────────────────────────
Shared style palette + cell-writing helpers for the section-driven Excel renderer.

The month-end tabs were one giant `render_excel` closure; the section bodies are
now extracted into `excel_sections.py`, and the shared state they all relied on
(colour palette, fonts, the `safe_text` / `write_currency` / `set_header_row`
helpers) lives here on `ExcelContext` so every section fn receives the SAME styling
and safety routing — behaviour-preserving, not a rewrite.

SAFETY (doc-rendering AP-DR-02 / AP-DR-05): every free-text cell routes through
`safe_text` (formula-injection neutralized + forced string type); every financial
figure routes through `write_currency` (a real number + accounting number_format),
never a pre-formatted currency string.

Typing note: `ws`/`wb` are untyped third-party openpyxl objects and the report
payload is provider-derived QBO/DB JSON of varying shape (read defensively via
.get) — both are genuine dynamic edges.
"""
from __future__ import annotations

from typing import Any, Dict, List  # CRAFT-OK: Any = openpyxl Worksheet/Workbook + dynamic QBO payload edge

# A worksheet/workbook is an opaque openpyxl object; Payload is the dynamic QBO
# report dict. Aliased so the dynamic edge is named once.
Worksheet = Any  # CRAFT-OK: opaque third-party openpyxl Worksheet boundary
Payload = Dict[str, Any]  # CRAFT-OK: dynamic provider-derived QBO report payload boundary


def neutralize_formula(value: object) -> str:
    """Apostrophe-prefix a string beginning with = + - @ (or a leading tab/CR) so
    a spreadsheet treats it as literal text, not an executable formula (OWASP
    CSV/formula injection). Every untrusted QBO/DB free-text value passes through
    this before it is written to a cell."""
    text = "" if value is None else str(value)
    if text and text[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text
    return text


def safe_float(val: object, default: float = 0.0) -> float:
    try:
        return float(val)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return default


def col_letter(col_num: int) -> str:
    """Convert 1-based column number to Excel letter."""
    result = ""
    while col_num > 0:
        col_num, remainder = divmod(col_num - 1, 26)
        result = chr(65 + remainder) + result
    return result


class ExcelContext:
    """Carries the workbook, palette, fonts, and the shared cell helpers that every
    section render fn uses. Constructed once per `render_excel` call. The openpyxl
    style classes are exposed as attributes (`Font`, `PatternFill`, `Alignment`) so
    sections build one-off styles without re-importing."""

    def __init__(self, wb: Worksheet, report_payload: Payload) -> None:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        # Expose the style classes for one-off styling inside sections.
        self.Font = Font
        self.PatternFill = PatternFill
        self.Alignment = Alignment

        self.wb = wb
        self.payload = report_payload

        # ── Colour palette (matching PDF) ──────────────────────────────────
        self.NAVY = "0F2744"
        self.NAVY_LIGHT = "1E3A5F"
        self.SILVER = "F0F4F8"
        self.WHITE = "FFFFFF"
        self.RED = "C0392B"
        self.GREEN = "16A34A"
        self.GRID = "D1D5DB"
        self.MUTED = "6B7280"
        self.TOTAL_BG = "E2E8F0"
        self.ACCENT = "1E40AF"

        # ── Styles ─────────────────────────────────────────────────────────
        self.header_font = Font(name="Calibri", bold=True, color=self.WHITE, size=10)
        self.header_fill = PatternFill(start_color=self.NAVY_LIGHT, end_color=self.NAVY_LIGHT, fill_type="solid")
        self.total_font = Font(name="Calibri", bold=True, size=10)
        self.total_fill = PatternFill(start_color=self.TOTAL_BG, end_color=self.TOTAL_BG, fill_type="solid")
        self.section_font = Font(name="Calibri", bold=True, size=10, color=self.NAVY)
        self.alt_fill = PatternFill(start_color=self.SILVER, end_color=self.SILVER, fill_type="solid")
        self.kpi_label_font = Font(name="Calibri", size=9, color=self.MUTED)
        self.kpi_val_font = Font(name="Calibri", bold=True, size=14, color=self.NAVY)
        self.title_font = Font(name="Calibri", bold=True, size=16, color=self.NAVY)
        self.subtitle_font = Font(name="Calibri", bold=True, size=12, color=self.NAVY)
        self.thin_border = Border(
            left=Side(style="thin", color=self.GRID),
            right=Side(style="thin", color=self.GRID),
            top=Side(style="thin", color=self.GRID),
            bottom=Side(style="thin", color=self.GRID),
        )
        self.acct_fmt_neg = '#,##0;[Red](#,##0);"-"'
        self.pct_fmt = '0.0%'

        # ── Bound metadata (display strings) ───────────────────────────────
        metadata = report_payload.get("metadata", {})
        self.company = neutralize_formula(metadata.get("company_name") or "Vesper Bot Report")
        self.period = metadata.get("period_label") or metadata.get("period", "—")
        self.gen_at = metadata.get("generated_at", "—")

    # ── Shared cell helpers (verbatim semantics from the old closure) ──────

    def safe_text(self, ws: Worksheet, row: int, col: int, value: object) -> Worksheet:
        cell = ws.cell(row=row, column=col, value=neutralize_formula(value))
        cell.data_type = "s"
        return cell

    def set_header_row(self, ws: Worksheet, row: int, headers: List[str], col_start: int = 1) -> None:
        for i, h in enumerate(headers, start=col_start):
            cell = ws.cell(row=row, column=i, value=h)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.Alignment(
                horizontal="center" if i > col_start else "left",
                vertical="center", wrap_text=True)
            cell.border = self.thin_border

    def write_currency(self, ws: Worksheet, row: int, col: int, value: object, bold: bool = False) -> None:
        cell = ws.cell(row=row, column=col, value=safe_float(value) if value is not None else 0)
        cell.number_format = self.acct_fmt_neg
        cell.alignment = self.Alignment(horizontal="right")
        cell.border = self.thin_border
        if bold:
            cell.font = self.total_font

    def write_pct(self, ws: Worksheet, row: int, col: int, value: object, bold: bool = False) -> None:
        cell = ws.cell(row=row, column=col, value=safe_float(value))
        cell.number_format = self.pct_fmt
        cell.alignment = self.Alignment(horizontal="right")
        cell.border = self.thin_border
        if bold:
            cell.font = self.total_font

    def write_number(self, ws: Worksheet, row: int, col: int, value: object,
                     decimals: int = 2, bold: bool = False) -> None:
        """Write a plain real NUMBER (a ratio/multiple like DSCR 1.25, NOT a currency)
        so it stays re-derivable in the cell (AP-DR-05). `None` renders an em-dash."""
        if value is None:
            cell = ws.cell(row=row, column=col, value="—")
        else:
            cell = ws.cell(row=row, column=col, value=safe_float(value))
            cell.number_format = "#,##0." + "0" * max(decimals, 1)
        cell.alignment = self.Alignment(horizontal="right")
        cell.border = self.thin_border
        if bold:
            cell.font = self.total_font

    def apply_alt_rows(self, ws: Worksheet, start_row: int, end_row: int, n_cols: int, col_start: int = 1) -> None:
        for r in range(start_row, end_row + 1):
            if (r - start_row) % 2 == 1:
                for c in range(col_start, col_start + n_cols):
                    ws.cell(row=r, column=c).fill = self.alt_fill

    def freeze_and_size(self, ws: Worksheet, freeze_cell: str, col_widths: Dict[int, float]) -> None:
        ws.freeze_panes = freeze_cell
        for col, width in col_widths.items():
            ws.column_dimensions[col_letter(col)].width = width
