"""
app/reporting/excel_renderer.py
────────────────────────────────
Section-driven board-pack Excel (.xlsx) renderer for Peregrine CFO reports.

`render_excel(payload, definition=None)` builds one `ExcelContext` and LOOPS the
definition's sections: for each, it looks up the `SectionType` in the registry,
checks its `condition`, runs its `binder`, and calls its `excel_render` fn. The
default definition is the month-end pack, which renders content-identically to the
old hardcoded renderer (proven by the golden in tests/reporting).

Mirrors the PDF report layout exactly (tab order = section order):
  Tab 1: P&L Metrics  ·  Tab 2: Profitability Trends  ·  Tab 3: P&L Detail
  Tab 4: Rolling 12M P&L  ·  Tab 5: Balance Sheet  ·  Tab 6: Rolling 12M CF
  Tab 7: BVA Variances (conditional)  ·  Tab 8: AR-AP Aging

TENANT-SAFETY PRECONDITION (doc-rendering AP-DR-07): the caller MUST gate tenant
access (check_tenant_access / _check_company_access) and build `report_payload`
from tenant-scoped queries BEFORE calling this renderer — a renderer is a pure
function of its payload and cannot enforce isolation. routers/reporting.py and
tasks/reporting.py satisfy this. SAFETY of free-text cells (AP-DR-02) +
numeric-currency cells (AP-DR-05) lives in ExcelContext + the section fns.
"""
from __future__ import annotations

import io
import logging
from typing import Any, Dict, Optional  # CRAFT-OK: Any = dynamic QBO payload boundary

from .report_definitions import MONTH_END_REPORT_DEFINITION, definition_sections
from .section_registry import SECTION_REGISTRY, ExcelContext

logger = logging.getLogger(__name__)

Payload = Dict[str, Any]  # CRAFT-OK: dynamic provider-derived QBO report payload boundary


def render_excel(report_payload: Payload, definition: Optional[Payload] = None) -> bytes:
    """Render a report payload into a multi-tab Excel workbook (.xlsx) by looping a
    report definition's sections (defaults to the month-end board pack)."""
    from openpyxl import Workbook

    if definition is None:
        definition = MONTH_END_REPORT_DEFINITION

    wb = Workbook()
    # openpyxl seeds one active sheet. The month-end pack's first section (kpis)
    # RENAMES + fills it; a definition whose first section instead `create_sheet`s its
    # own tab (e.g. standard_financials → income_statement) leaves this one empty. We
    # capture it here and drop it after rendering IF it stayed empty and other tabs
    # exist — so no definition emits a stray blank leading tab.
    default_sheet = wb.active
    ctx = ExcelContext(wb, report_payload)
    rendered_any = False

    for spec in definition_sections(definition):
        section_id = spec.get("section_id")  # SA-OK: bounded-loop over the definition's fixed section list, not a DB query
        section = SECTION_REGISTRY.get(section_id) if section_id else None
        if section is None:
            continue
        if not section.condition(report_payload):
            continue
        section_data = section.binder(report_payload)
        section.excel_render(ctx, section_data)
        rendered_any = True

    if not rendered_any:
        # A definition rendered NOTHING — leave a single named sheet so wb.save
        # succeeds rather than raising on an empty workbook.
        default_sheet.title = "Report"
    elif default_sheet.max_row == 1 and default_sheet.max_column == 1 \
            and default_sheet["A1"].value is None and len(wb.worksheets) > 1:
        # The default sheet was never written to (a section created its own tab first)
        # and real tabs exist — remove the stray blank leading sheet.
        wb.remove(default_sheet)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
